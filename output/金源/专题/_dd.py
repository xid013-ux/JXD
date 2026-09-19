# -*- coding: utf-8 -*-
"""尽职调查报告体例的 Word 构建件。

参数取自 output/巨蟹/巨蟹智能尽调报告第六章_公司所处行业分析.docx 实测值，
与任务说明第八节的模板参数一致：

  一级标题  黑体 14pt 不加粗，无缩进，行距 1.5，段前段后各 6pt
  二级标题  宋体加粗 12pt，首行缩进 2 字符，段前 6pt
  正文      宋体 12pt／西文 Times New Roman，首行缩进 2 字符，行距 1.5，两端对齐
  表题图题  宋体加粗 10.5pt，居中，置于表／图上方
  资料来源  宋体 10.5pt，无缩进，置于表／图下方，单独成行
  表格      宽度占满版心、居中、固定列宽；外框线粗内线细；首行加粗居中并跨页重复
"""
import os
from docx import Document
from docx.shared import Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

_HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.environ.get('DD_TPL',
                     'output/巨蟹/巨蟹智能尽调报告第六章_公司所处行业分析.docx')
OUT = os.environ.get('DD_OUT', os.path.join(os.getcwd(), 'out.docx'))
CH = os.environ.get('DD_CHARTS', os.path.join(os.getcwd(), 'charts')) + os.sep

doc = Document(TPL)
body = doc.element.body
sectPr = body.find(qn('w:sectPr'))
for child in list(body):
    if child is not sectPr:
        body.remove(child)

_pg = sectPr.find(qn('w:pgSz'))
_mar = sectPr.find(qn('w:pgMar'))
CONTENT_TW = (int(_pg.get(qn('w:w')))
              - int(_mar.get(qn('w:left'))) - int(_mar.get(qn('w:right'))))

_CN = {'1': '一', '2': '二', '3': '三', '4': '四', '5': '五',
       '6': '六', '7': '七', '8': '八', '9': '九'}
_h1n = [0]
_h2n = [0]


def _q(t):
    """ASCII 直引号成对转为中文弯引号。"""
    if not isinstance(t, str) or '"' not in t:
        return t
    out, open_ = [], True
    for ch in t:
        if ch == '"':
            out.append('“' if open_ else '”')
            open_ = not open_
        else:
            out.append(ch)
    return ''.join(out)


def _p(text, *, ea='宋体', sz=24, bold=False, indent=None, jc='both',
       line=360, before=120, after=None, keep_next=False):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    # CT_PPr 规定 keepNext 在 spacing / ind / jc 之前
    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))
    sp = OxmlElement('w:spacing')
    if line:
        sp.set(qn('w:line'), str(line))
        sp.set(qn('w:lineRule'), 'auto')
    sp.set(qn('w:before'), str(before))
    if after is not None:
        sp.set(qn('w:after'), str(after))
    pPr.append(sp)
    if indent:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), str(indent))
        pPr.append(ind)
    j = OxmlElement('w:jc')
    j.set(qn('w:val'), jc)
    pPr.append(j)
    p.append(pPr)
    if text:
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        f = OxmlElement('w:rFonts')
        f.set(qn('w:ascii'), 'Times New Roman')
        f.set(qn('w:hAnsi'), 'Times New Roman')
        f.set(qn('w:eastAsia'), ea)
        rPr.append(f)
        b = OxmlElement('w:b')
        if not bold:
            b.set(qn('w:val'), '0')
        rPr.append(b)
        s = OxmlElement('w:sz')
        s.set(qn('w:val'), str(sz))
        rPr.append(s)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = _q(text)
        r.append(t)
        p.append(r)
    body.insert(list(body).index(sectPr), p)
    return p


def h1(text):
    """一级标题：一、二、三……自动编号。"""
    _h1n[0] += 1
    _h2n[0] = 0
    return _p(f'{_CN[str(_h1n[0])]}、{text}', ea='黑体', sz=28,
              bold=False, indent=None, after=120)


def h2(text):
    """二级标题：（一）（二）……自动编号。"""
    _h2n[0] += 1
    return _p(f'（{_CN[str(_h2n[0])]}）{text}', ea='宋体', sz=24,
              bold=True, indent=482)


def para(text):
    return _p(text, indent=480)


def cap(text):
    """表题／图题，置于表图上方。"""
    return _p(text, sz=21, bold=True, indent=None, jc='center',
              line=None, before=120, after=60, keep_next=True)


def src(text):
    """资料来源，置于表图下方。"""
    return _p(text, sz=21, indent=None, jc='both',
              line=None, before=0, after=120)


def _borders(tblPr):
    bd = OxmlElement('w:tblBorders')
    for tag, sz in (('top', 12), ('left', 12), ('bottom', 12), ('right', 12),
                    ('insideH', 4), ('insideV', 4)):
        e = OxmlElement(f'w:{tag}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '010000')
        bd.append(e)
    tblPr.append(bd)


def _cell(text, *, bold, align, width):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    w = OxmlElement('w:tcW')
    w.set(qn('w:w'), str(width))
    w.set(qn('w:type'), 'dxa')
    tcPr.append(w)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'center')
    tcPr.append(va)
    tc.append(tcPr)
    for line in str(text).split('\n'):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '40')
        sp.set(qn('w:after'), '40')
        sp.set(qn('w:line'), '260')
        sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
        j = OxmlElement('w:jc')
        j.set(qn('w:val'), align)
        pPr.append(j)
        p.append(pPr)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        f = OxmlElement('w:rFonts')
        f.set(qn('w:ascii'), 'Times New Roman')
        f.set(qn('w:hAnsi'), 'Times New Roman')
        f.set(qn('w:eastAsia'), '宋体')
        rPr.append(f)
        b = OxmlElement('w:b')
        if not bold:
            b.set(qn('w:val'), '0')
        rPr.append(b)
        s = OxmlElement('w:sz')
        s.set(qn('w:val'), '21')
        rPr.append(s)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = _q(line)
        r.append(t)
        p.append(r)
        tc.append(p)
    return tc


def table(headers, rows, widths, caption, source, aligns=None):
    """固定列宽表格。widths 为权重，按版心宽度归一化为 dxa。"""
    cap(caption)
    aligns = aligns or ['center'] * len(headers)
    tot = sum(widths)
    cols = [int(CONTENT_TW * w / tot) for w in widths]
    cols[-1] += CONTENT_TW - sum(cols)

    tbl = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), '5000')
    w.set(qn('w:type'), 'pct')
    tblPr.append(w)
    j = OxmlElement('w:jc')
    j.set(qn('w:val'), 'center')
    tblPr.append(j)
    # CT_TblPrBase 规定 tblBorders 在 tblLayout 之前
    _borders(tblPr)
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    tbl.append(tblPr)

    grid = OxmlElement('w:tblGrid')
    for c in cols:
        g = OxmlElement('w:gridCol')
        g.set(qn('w:w'), str(c))
        grid.append(g)
    tbl.append(grid)

    for ri, row in enumerate([headers] + list(rows)):
        tr = OxmlElement('w:tr')
        trPr = OxmlElement('w:trPr')
        trPr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trPr.append(OxmlElement('w:tblHeader'))
        jj = OxmlElement('w:jc')
        jj.set(qn('w:val'), 'center')
        trPr.append(jj)
        tr.append(trPr)
        for ci, val in enumerate(row):
            tr.append(_cell(val, bold=(ri == 0),
                            align='center' if ri == 0 else aligns[ci],
                            width=cols[ci]))
        tbl.append(tr)

    body.insert(list(body).index(sectPr), tbl)
    src(source)
    return tbl


def figure(fname, caption, source):
    """图：图题在上，图居中，资料来源在下。"""
    cap(caption)
    path = CH + fname
    iw, ih = Image.open(path).size
    width = Emu(int(CONTENT_TW / 1440 * 914400))
    height = Emu(int(width * ih / iw))
    p = _p('', indent=None, jc='center', line=None, before=0, after=0,
           keep_next=True)
    run = doc.add_paragraph().add_run()
    run.add_picture(path, width=width, height=height)
    holder = run._element.getparent()
    p.append(run._element)
    holder.getparent().remove(holder)
    src(source)
    return p


def set_header(company, doctype):
    """改写页眉。模板继承自既往项目，页眉里写死着上一家公司的名称，
    不改写会把别家客户的名字带进交付物。"""
    _HDR['company'] = company
    _HDR['doctype'] = doctype


_HDR = {'company': None, 'doctype': None}
# 模板沿用自既往项目，页眉中残留的公司名称
_STALE = ['无锡巨蟹智能驱动科技有限公司', '江苏耀鸿电子', '有限公司']


def _rewrite_headers(path):
    import re, shutil, zipfile
    if not _HDR['company']:
        return
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path)
    hit = 0
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r'word/header\d+\.xml$', item.filename):
                x = data.decode('utf-8')
                texts = re.findall(r'<w:t[^>]*>([^<]*)</w:t>', x)
                if any(s in ''.join(texts) for s in _STALE):
                    hit += 1
                    first = True
                    def repl(m):
                        nonlocal first
                        t = m.group(1)
                        if first:
                            first = False
                            return m.group(0).replace(
                                '>' + t + '<', '>' + _HDR['company'] + '<')
                        if '尽职调查报告' in t:
                            return m.group(0).replace(
                                '尽职调查报告', _HDR['doctype'])
                        if any(s in t for s in _STALE):
                            return m.group(0).replace('>' + t + '<', '><')
                        return m.group(0)
                    x = re.sub(r'<w:t[^>]*>([^<]*)</w:t>', repl, x)
                    data = x.encode('utf-8')
            zout.writestr(item, data)
    zin.close()
    shutil.move(tmp, path)
    # 复核：交付物中不得残留其他公司名称
    zc = zipfile.ZipFile(path)
    for item in zc.infolist():
        if item.filename.endswith('.xml'):
            body_txt = zc.read(item.filename).decode('utf-8', 'ignore')
            for s in _STALE:
                if s in body_txt and s != '有限公司':
                    raise AssertionError(f'{item.filename} 仍残留 {s}')
    zc.close()
    print(f'页眉已改写（{hit} 处）')


def save():
    doc.save(OUT)
    _rewrite_headers(OUT)
    print('saved', OUT)
